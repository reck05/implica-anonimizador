"""Procesador de PowerPoint (.pptx) con python-pptx.

Procesa: text frames en shapes, tablas, notas de speaker, hipervínculos y
metadatos del documento (autor/título/etc.). Además, un barrido XML a nivel ZIP
cubre TODO el texto que python-pptx no expone bien: GRÁFICOS (títulos/categorías/
series), SmartArt/diagramas, patrones/plantillas y cuadros de texto.
Preserva formato editando .text de cada run.
"""
from __future__ import annotations

import glob
import os
import shutil
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation

from ..replacer import Replacer
from ._meta import extract_core_props, scrub_core_props

# Namespaces OOXML: texto DrawingML (<a:t>) y valores de gráfico (<c:v>).
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
# Subcarpetas de ppt/ donde vive texto que python-pptx no toca bien.
_SWEEP_SUBDIRS = ("slides", "notesSlides", "charts", "diagrams",
                  "slideMasters", "slideLayouts")


def _iter_text_frames(shape):
    """Yield text_frames recursivos (groups, tables, etc.)."""
    if shape.has_text_frame:
        yield shape.text_frame
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame
    if shape.shape_type == 6:  # GROUP
        for sub in shape.shapes:
            yield from _iter_text_frames(sub)


def _replace_in_text_frame(tf, replacer: Replacer) -> int:
    total = 0
    for paragraph in tf.paragraphs:
        full = "".join(r.text for r in paragraph.runs)
        new, n = replacer.apply(full)
        if n == 0:
            # Aun sin cambios en el texto visible, puede haber un hipervínculo
            # con un nombre en la URL.
            _replace_run_hyperlinks(paragraph, replacer)
            continue
        total += n
        # Intento por run
        runs_changed = 0
        for run in paragraph.runs:
            run_new, run_n = replacer.apply(run.text)
            if run_n > 0:
                run.text = run_new
                runs_changed += run_n
            _replace_run_hyperlinks_single(run, replacer)
        if runs_changed < n and paragraph.runs:
            paragraph.runs[0].text = new
            for r in paragraph.runs[1:]:
                r.text = ""
    return total


def _replace_run_hyperlinks_single(run, replacer: Replacer) -> None:
    """Reemplaza en la URL del hipervínculo de un run (p.ej. ?cliente=GlobalMenta)."""
    try:
        hl = run.hyperlink
        if hl is not None and hl.address:
            new_addr, n = replacer.apply(hl.address)
            if n > 0:
                hl.address = new_addr
    except Exception:
        pass


def _replace_run_hyperlinks(paragraph, replacer: Replacer) -> None:
    for run in paragraph.runs:
        _replace_run_hyperlinks_single(run, replacer)


def count_images(path: Path) -> int:
    """Cuenta imágenes embebidas (los nombres DENTRO de una imagen no se anonimizan)."""
    count = 0
    try:
        prs = Presentation(str(path))
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.shape_type == 13:  # PICTURE
                    count += 1
    except Exception:
        return 0
    return count


def _extract_pptx_xml_texts(path: Path) -> list[str]:
    """Texto en partes que python-pptx no expone: gráficos, SmartArt, plantillas."""
    out: list[str] = []
    try:
        from lxml import etree
        with zipfile.ZipFile(str(path)) as z:
            wanted = [n for n in z.namelist()
                      if any(n.startswith(f"ppt/{sub}/") for sub in _SWEEP_SUBDIRS)
                      and n.endswith(".xml")]
            for n in wanted:
                try:
                    tree = etree.fromstring(z.read(n))
                    for tag in ("{%s}t" % _A_NS, "{%s}v" % _C_NS):
                        for el in tree.iter(tag):
                            if el.text and el.text.strip():
                                out.append(el.text)
                except Exception:
                    continue
    except Exception:
        pass
    return out


def extract_text(path: Path) -> list[str]:
    prs = Presentation(str(path))
    texts: list[str] = []
    texts.extend(extract_core_props(prs.core_properties, "core"))
    for slide in prs.slides:
        for shape in slide.shapes:
            for tf in _iter_text_frames(shape):
                for p in tf.paragraphs:
                    line = "".join(r.text for r in p.runs)
                    if line.strip():
                        texts.append(line)
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            for p in notes_tf.paragraphs:
                line = "".join(r.text for r in p.runs)
                if line.strip():
                    texts.append(line)
    texts.extend(_extract_pptx_xml_texts(path))
    return texts


def _sweep_pptx_xml(path: Path, replacer: Replacer) -> None:
    """Reemplaza en TODO el texto DrawingML (<a:t>) y valores de gráfico (<c:v>)
    de las partes que python-pptx no cubre (gráficos, SmartArt, patrones, cuadros
    de texto). Idempotente sobre lo ya reemplazado y DEFENSIVO: ante cualquier
    error deja el .pptx de python-pptx intacto (no corrompe)."""
    tmp = None
    try:
        from lxml import etree
        tmp = tempfile.mkdtemp()
        with zipfile.ZipFile(str(path)) as z:
            z.extractall(tmp)
        pptdir = os.path.join(tmp, "ppt")
        if not os.path.isdir(pptdir):
            return
        xml_files: list[str] = []
        for sub in _SWEEP_SUBDIRS:
            xml_files += glob.glob(os.path.join(pptdir, sub, "*.xml"))
        tags = ("{%s}t" % _A_NS, "{%s}v" % _C_NS)
        for xf in xml_files:
            try:
                tree = etree.parse(xf)
                changed = False
                for tag in tags:
                    for el in tree.iter(tag):
                        if el.text:
                            new, n = replacer.apply(el.text)
                            if n > 0:
                                el.text = new
                                changed = True
                if changed:
                    tree.write(xf, xml_declaration=True, encoding="UTF-8", standalone=True)
            except Exception:
                continue
        tmp_zip = str(path) + ".tmp"
        with zipfile.ZipFile(tmp_zip, "w", zipfile.ZIP_DEFLATED) as z:
            for root, _dirs, files in os.walk(tmp):
                for f in files:
                    fp = os.path.join(root, f)
                    arc = os.path.relpath(fp, tmp).replace(os.sep, "/")
                    z.write(fp, arc)
        shutil.move(tmp_zip, str(path))
    except Exception:
        pass
    finally:
        if tmp:
            try:
                shutil.rmtree(tmp)
            except Exception:
                pass


def apply_replacements(src: Path, mapping: dict[str, str], dst: Path) -> None:
    replacer = Replacer(mapping)
    prs = Presentation(str(src))
    for slide in prs.slides:
        for shape in slide.shapes:
            for tf in _iter_text_frames(shape):
                _replace_in_text_frame(tf, replacer)
        if slide.has_notes_slide:
            _replace_in_text_frame(slide.notes_slide.notes_text_frame, replacer)
    scrub_core_props(prs.core_properties, replacer, "core")
    prs.save(str(dst))
    # Gráficos, SmartArt, patrones y cuadros de texto que python-pptx no cubre.
    _sweep_pptx_xml(dst, replacer)
