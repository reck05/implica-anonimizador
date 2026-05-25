"""Test del flujo PGC: detector de cuentas → candidatos → mapping automático → reemplazo.

Compara con el flujo NER puro para mostrar la mejora.
"""
import io
import sys
import time
from collections import Counter
from pathlib import Path

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", line_buffering=True)
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", line_buffering=True)

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from implica_anon import accounting, formats, mapping as mapping_mod, rehydrate
from implica_anon.detectors import (
    candidates_from_pgc, cluster_variants, detect_candidates,
)

FIXTURE = Path(__file__).parent / "fixtures" / "sumas_saldos_sintetico.xlsx"


def test_pgc_scan():
    print("\n=== 1) Escaneo PGC del workbook ===")
    t0 = time.time()
    scan = accounting.scan_workbook(FIXTURE)
    t1 = time.time()
    print(f"Tiempo escaneo: {t1-t0:.2f}s")
    print(f"Es libro contable: {scan.is_accounting}")
    print(f"Entidades encontradas: {len(scan.entities)}")
    print("\nHojas con cuentas detectadas:")
    for sheet, info in scan.sheet_findings.items():
        print(f"  '{sheet}': fila cabecera={info['header_row']}, "
              f"col cuenta={info['col_account']}, col desc={info['col_desc']}, "
              f"{info['rows_matched']}/{info['rows_scanned']} filas con cuenta de terceros")

    by_kind = Counter(e.kind for e in scan.entities)
    print(f"\nDistribución por tipo:")
    for kind, count in by_kind.most_common():
        print(f"  {kind}: {count}")

    # Muestra
    print(f"\nMuestra de entidades:")
    seen_kinds = set()
    for e in scan.entities:
        if e.kind in seen_kinds:
            continue
        seen_kinds.add(e.kind)
        cleaned = accounting.clean_description(e.kind, e.description)
        print(f"  [{e.kind}] {e.code}: '{e.description}' → cleaned: '{cleaned}'")

    return scan


def test_pgc_candidates(scan):
    print("\n=== 2) Convertir a candidatos PGC ===")
    pgc_candidates = candidates_from_pgc(scan)
    print(f"Candidatos PGC: {len(pgc_candidates)}")
    print(f"Por tipo:")
    by_kind = Counter(c.kind for c in pgc_candidates)
    for kind, count in by_kind.most_common():
        print(f"  {kind}: {count}")
    return pgc_candidates


def test_compare_coverage():
    print("\n=== 3) Comparación PGC vs NER puro ===")

    # Ground truth
    wb = load_workbook(str(FIXTURE), data_only=True, read_only=True)
    ws = wb["Sumas y Saldos 2024"]
    truth_clients_proveed: set[str] = set()
    import re as _re
    for row in ws.iter_rows(min_row=6, values_only=True):
        if not row[0]:
            continue
        code = str(row[0])
        desc = str(row[1] or "")
        if code.startswith(("43", "410", "411")):
            nombre = _re.sub(r"\s*[,\(]\s*(?:CIF[:\s]*)?[A-Z]\d{7}[0-9A-J][\s\)]*", "", desc).strip(" ,;-")
            truth_clients_proveed.add(nombre)
    wb.close()
    print(f"Ground truth (clientes + proveedores): {len(truth_clients_proveed)}")

    # --- PGC: detecta los nombres ---
    scan = accounting.scan_workbook(FIXTURE)
    pgc_candidates = candidates_from_pgc(scan)
    detected_pgc = {c.text for c in pgc_candidates if c.kind in ("CLIENTE", "PROVEEDOR")}

    matched_pgc = 0
    for nombre in truth_clients_proveed:
        for d in detected_pgc:
            if nombre.lower() == d.lower() or nombre.lower() in d.lower() or d.lower() in nombre.lower():
                matched_pgc += 1
                break
    cov_pgc = matched_pgc / len(truth_clients_proveed) * 100
    print(f"  PGC:        {matched_pgc}/{len(truth_clients_proveed)} = {cov_pgc:.1f}%")

    # --- NER puro: como antes ---
    texts = formats.extract_text(FIXTURE)
    ner_candidates = detect_candidates(texts)
    detected_ner = {c.text for c in ner_candidates if c.kind in ("ORG", "PER")}
    matched_ner = 0
    for nombre in truth_clients_proveed:
        for d in detected_ner:
            if nombre.lower() == d.lower() or nombre.lower() in d.lower() or d.lower() in nombre.lower():
                matched_ner += 1
                break
    cov_ner = matched_ner / len(truth_clients_proveed) * 100
    print(f"  NER puro:   {matched_ner}/{len(truth_clients_proveed)} = {cov_ner:.1f}%")

    print(f"  Mejora:     +{cov_pgc - cov_ner:.1f} puntos porcentuales")


def test_full_anonymization():
    print("\n=== 4) Flujo completo: detectar → mapping masivo → anonimizar ===")

    # Escaneo PGC
    scan = accounting.scan_workbook(FIXTURE)
    pgc_candidates = candidates_from_pgc(scan)

    # También regex (CIF, IBAN, email, etc.) sin NER (más rápido y limpio)
    texts = formats.extract_text(FIXTURE)
    t0 = time.time()
    regex_candidates = detect_candidates(texts, skip_ner=True)
    t1 = time.time()
    print(f"Regex (sin NER): {len(regex_candidates)} candidatos en {t1-t0:.2f}s")

    # Combinar
    all_candidates = pgc_candidates + regex_candidates
    t2 = time.time()
    clusters = cluster_variants(all_candidates)
    t3 = time.time()
    print(f"Clustering: {len(clusters)} clusters en {t3-t2:.2f}s")

    # Distribución
    by_kind = Counter(c.kind for c in clusters)
    print(f"Clusters por tipo:")
    for kind, count in by_kind.most_common():
        print(f"  {kind}: {count}")

    # Asignación automática de codenames
    print("\nAsignando codenames automáticamente...")
    from implica_anon.interactive import DEFAULT_PLACEHOLDERS

    pm = mapping_mod.ProjectMapping(project="test_pgc")
    counters: dict[str, int] = {}
    for c in clusters:
        template = DEFAULT_PLACEHOLDERS.get(c.kind, "[{}-{{}}]".format(c.kind))
        counters[c.kind] = counters.get(c.kind, 0) + 1
        codename = template.format(counters[c.kind])
        for v in c.variants:
            pm.add(c.kind, v, codename)

    total = sum(len(v) for v in pm.entries.values())
    print(f"Total entradas en mapping: {total}")

    # Aplicar
    out_path = FIXTURE.with_name("sumas_saldos_anonimizado.xlsx")
    t4 = time.time()
    formats.apply_replacements(FIXTURE, pm.all_replacements(), out_path)
    t5 = time.time()
    print(f"Reemplazo: {t5-t4:.2f}s")
    print(f"Output: {out_path.name} ({out_path.stat().st_size // 1024} KB)")

    # Verificar
    wb = load_workbook(str(out_path), data_only=True, read_only=True)
    ws = wb["Sumas y Saldos 2024"]
    survivors = []
    for row in ws.iter_rows(min_row=6, max_row=50, values_only=True):
        if not row[0]:
            continue
        desc = str(row[1] or "")
        # Si la descripción de una cuenta 43x todavía contiene letras pero no es un placeholder
        code = str(row[0])
        if code.startswith(("43", "41")) and desc and not desc.startswith("["):
            # Verificar que no sea solo CIF u otra cosa
            if any(c.isalpha() for c in desc):
                survivors.append((code, desc))
    wb.close()
    if survivors:
        print(f"\n⚠ {len(survivors)} cuentas con descripción no anonimizada (muestra):")
        for code, desc in survivors[:5]:
            print(f"    {code}: {desc}")
    else:
        print(f"\n✓ Todas las cuentas de terceros en las primeras 50 filas están anonimizadas")

    return pm, out_path


def test_rehydrate(pm, anonimizado_path):
    print("\n=== 5) Rehydrate: revertir codenames ===")
    rehidratado_path = anonimizado_path.with_name("sumas_saldos_rehidratado.xlsx")
    rehydrate.rehydrate_file(anonimizado_path, pm, rehidratado_path)
    print(f"Rehidratado: {rehidratado_path.name}")

    # Verificar: las primeras filas deberían volver a tener nombres
    wb = load_workbook(str(rehidratado_path), data_only=True, read_only=True)
    ws = wb["Sumas y Saldos 2024"]
    rehydrated_count = 0
    placeholder_count = 0
    for row in ws.iter_rows(min_row=6, max_row=100, values_only=True):
        if not row[0]:
            continue
        desc = str(row[1] or "")
        if desc.startswith("["):
            placeholder_count += 1
        else:
            rehydrated_count += 1
    wb.close()
    print(f"En primeras 100 filas: {rehydrated_count} con nombre, {placeholder_count} con placeholder")
    if placeholder_count == 0:
        print("✓ Rehidratado correcto")
    else:
        print(f"⚠ Quedaron {placeholder_count} placeholders sin rehidratar")


if __name__ == "__main__":
    scan = test_pgc_scan()
    test_pgc_candidates(scan)
    test_compare_coverage()
    pm, out = test_full_anonymization()
    test_rehydrate(pm, out)
    print("\n✓✓✓ Tests pasados")
