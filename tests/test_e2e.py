"""Test end-to-end con datos sintéticos.

Verifica:
1. Detección de entidades (NER + regex)
2. Clustering de variantes
3. Reemplazo en Excel preservando estructura
"""
import io
import sys
from pathlib import Path

# Forzar UTF-8 en stdout para evitar problemas con cp1252 en PowerShell
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", line_buffering=True)
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", line_buffering=True)

# Permite ejecutar desde la raíz del repo
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from implica_anon import formats, mapping as mapping_mod
from implica_anon.detectors import cluster_variants, detect_candidates

FIXTURE = Path(__file__).parent / "fixtures" / "deal_sintetico.xlsx"


def test_detection():
    print("\n=== 1) Extracción de texto ===")
    texts = formats.extract_text(FIXTURE)
    print(f"  Extraídos {len(texts)} fragmentos de texto")
    for t in texts[:5]:
        print(f"    > {t}")

    print("\n=== 2) Detección de candidatos ===")
    candidates = detect_candidates(texts)
    print(f"  {len(candidates)} candidatos detectados")
    by_kind = {}
    for c in candidates:
        by_kind.setdefault(c.kind, []).append(c.text)
    for kind, items in by_kind.items():
        print(f"    [{kind}] {items}")

    print("\n=== 3) Clustering ===")
    clusters = cluster_variants(candidates)
    print(f"  {len(clusters)} clusters")
    for cl in clusters:
        print(f"    [{cl.kind}] {cl.canonical}  variantes={cl.variants}  count={cl.total_count}")

    # Verificaciones mínimas
    org_clusters = [c for c in clusters if c.kind == "ORG"]
    assert len(org_clusters) >= 2, f"Esperaba >=2 empresas, encontré {len(org_clusters)}"
    print("  ✓ Detectó al menos 2 empresas")

    # Que las variantes de Global Menta estén en el mismo cluster
    gm_cluster = next(
        (c for c in org_clusters if any("Menta" in v for v in c.variants)),
        None,
    )
    assert gm_cluster is not None, "No encontró cluster de Global Menta"
    print(f"  Variantes Global Menta: {gm_cluster.variants}")
    has_full = any("S.L." in v or "SL" in v for v in gm_cluster.variants)
    has_short = any("GlobalMenta" == v for v in gm_cluster.variants)
    assert has_short, (
        f"'GlobalMenta' (sin espacio) no clusterizado con 'Global Menta': {gm_cluster.variants}"
    )
    assert has_full, (
        f"No se expandió 'Global Menta' a 'Global Menta S.L.': {gm_cluster.variants}"
    )
    print("  ✓ Variantes 'Global Menta', 'GlobalMenta' y forma con sufijo S.L. clusterizadas")

    cif_clusters = [c for c in clusters if c.kind == "CIF"]
    assert len(cif_clusters) >= 2, f"Esperaba >=2 CIFs, encontré {len(cif_clusters)}"
    print(f"  ✓ Detectó {len(cif_clusters)} CIFs")

    iban_clusters = [c for c in clusters if c.kind == "IBAN"]
    assert len(iban_clusters) >= 1, "Esperaba >=1 IBAN"
    print(f"  ✓ Detectó {len(iban_clusters)} IBANs")

    return clusters


def test_replacement(clusters):
    print("\n=== 4) Mapping pre-cargado + reemplazo ===")
    pm = mapping_mod.ProjectMapping(project="paradise_test")

    # Asignar codenames a mano (simula resolve_clusters)
    for cl in clusters:
        if cl.kind == "ORG":
            codename = "Paradise" if "Menta" in cl.canonical else "Eden"
            for v in cl.variants:
                pm.add("ORG", v, codename)
        elif cl.kind == "PER":
            for v in cl.variants:
                pm.add("PER", v, f"[Contact-{v[:1]}]")
        elif cl.kind == "CIF":
            idx = len(pm.entries.get("CIF", {})) + 1
            for v in cl.variants:
                pm.add("CIF", v, f"[CIF-{idx}]")
        elif cl.kind == "IBAN":
            idx = len(pm.entries.get("IBAN", {})) + 1
            for v in cl.variants:
                pm.add("IBAN", v, f"[IBAN-{idx}]")
        elif cl.kind == "EMAIL":
            idx = len(pm.entries.get("EMAIL", {})) + 1
            for v in cl.variants:
                pm.add("EMAIL", v, f"[email-{idx}]")
        elif cl.kind == "ADDRESS":
            idx = len(pm.entries.get("ADDRESS", {})) + 1
            for v in cl.variants:
                pm.add("ADDRESS", v, f"[Dirección-{idx}]")

    print("  Mapping construido:")
    for kind, entries in pm.entries.items():
        for k, v in entries.items():
            print(f"    [{kind}] '{k}' -> '{v}'")

    # Aplicar
    out_path = FIXTURE.with_name("deal_sintetico.paradise.xlsx")
    formats.apply_replacements(FIXTURE, pm.all_replacements(), out_path)
    print(f"\n  Output: {out_path}")

    # Verificar contenido del output
    print("\n=== 5) Verificación del output ===")
    wb = load_workbook(str(out_path))
    found_paradise = False
    found_eden = False
    found_original = False
    for ws in wb.worksheets:
        print(f"  Hoja: '{ws.title}'")
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is None:
                    continue
                if isinstance(cell, str):
                    if "Paradise" in cell:
                        found_paradise = True
                    if "Eden" in cell:
                        found_eden = True
                    if "Global Menta" in cell or "Tropical Frutas" in cell or "GlobalMenta" in cell:
                        found_original = True
                        print(f"    ✗ Nombre original sobrevivió: {cell}")

    assert found_paradise, "No se encontró 'Paradise' en el output"
    assert found_eden, "No se encontró 'Eden' en el output"
    assert not found_original, "Quedaron nombres originales sin reemplazar"
    print("  ✓ 'Paradise' presente en el output")
    print("  ✓ 'Eden' presente en el output")
    print("  ✓ Ningún nombre original sobrevivió")

    # Verificar que los títulos de hoja se renombraron
    sheet_names = [ws.title for ws in wb.worksheets]
    print(f"  Hojas: {sheet_names}")
    assert any("Paradise" in s for s in sheet_names), (
        "El título de hoja con 'Global Menta' no se renombró"
    )
    print("  ✓ Nombre de hoja anonimizado")


if __name__ == "__main__":
    clusters = test_detection()
    test_replacement(clusters)
    print("\n✓✓✓ Todos los tests pasaron")
