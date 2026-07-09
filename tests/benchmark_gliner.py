"""Benchmark LOCAL: motor clásico (spaCy+regex+heurísticas) vs +GLiNER.

Genera 3 documentos ficticios M&A en español con entidades CONOCIDAS (ground truth)
y compara ambos motores. Vuelca un informe orientado a DECISIÓN DE EQUIPO en:
  - benchmark_resultado.md   (informe legible para jefe/IT)
  - benchmark_resultado.csv  (para Excel)
  - benchmark_resultado.json (trazabilidad)

REQUISITOS para números reales:
  - spaCy con es_core_news_md (motor clásico).
  - (opcional) GLiNER para la columna CON GLiNER:
        pip install -r requirements-gliner.txt
        python -m implica_anon.gliner_detector        # descarga el modelo una vez

USO:
  python tests/benchmark_gliner.py                    # solo clásico
  # Con GLiNER (PowerShell):
  $env:IMPLICA_ENABLE_GLINER="true"; $env:IMPLICA_GLINER_ALLOW_DOWNLOAD="true"
  python tests/benchmark_gliner.py

NOTA: si spaCy está bloqueado/no instalado, el motor clásico cae a heurísticas
(menos preciso) y el informe lo indica — los números solo son representativos en una
máquina con spaCy operativo.
"""
import csv
import io
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", line_buffering=True)
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", line_buffering=True)

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

ROOT = Path(__file__).resolve().parents[1]
# Carpeta única de salida: fixtures ficticios + informes (html/md/csv/json) juntos.
OUT = ROOT / "benchmark_outputs"
OUT.mkdir(parents=True, exist_ok=True)
FIX = OUT  # los documentos ficticios se generan en la misma carpeta de salida

GROUND_TRUTH = {
    "teaser": [
        "Proyecto Halcón", "Clínica Dental Sonrisa S.L.", "Implica Corporate Finance",
        "Mediterráneo Capital Partners", "DentalGroup Europe", "Alberto Ramírez",
        "María Fernández", "sonrisadental.es",
    ],
    "excel": [
        "Distribuciones García S.L.", "Mercadona S.A.", "Supermercados Día",
        "Suministros Dentales Ibéricos S.L.", "Material Médico Europa S.A.",
        "Banco Santander", "CaixaBank",
        "Laura Gómez Ruiz", "Javier Torres Molina",
        "Inversiones Halcón S.L.", "Patrimonial Ramírez S.A.",
    ],
    "pptx": [
        "Proyecto Halcón", "Salud Capital Partners", "Grupo Ramírez S.L.",
        "Implica Corporate Finance", "Mediterráneo Capital Partners",
        "Garrigues", "Alberto Ramírez",
    ],
}

DOC_LABEL = {"teaser": "Teaser M&A (PDF)", "excel": "Excel listado (no contable)",
             "pptx": "Presentación (PPTX)"}


def make_teaser_pdf() -> Path:
    import fitz
    path = FIX / "teaser_proyecto_halcon.pdf"
    lines = [
        "PROYECTO HALCÓN — Información Confidencial", "",
        "Implica Corporate Finance presenta una oportunidad de inversión en el sector dental.", "",
        "La compañía objetivo, Clínica Dental Sonrisa S.L. (CIF B12345678), es una cadena",
        "de clínicas dentales con sede en Valencia y web sonrisadental.es.", "",
        "Fundada por el Dr. Alberto Ramírez (CEO), la dirección financiera está a cargo de",
        "María Fernández (CFO). La compañía ha mostrado un crecimiento sólido.", "",
        "Entre los inversores interesados figura Mediterráneo Capital Partners, fondo de",
        "private equity especializado en salud. Como comparable de mercado se cita a",
        "DentalGroup Europe.", "",
        "Para más información, contactar con el asesor financiero Implica Corporate Finance.",
    ]
    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for ln in lines:
        page.insert_text((72, y), ln, fontsize=11, fontname="helv")
        y += 18
    doc.save(str(path))
    doc.close()
    return path


def make_excel() -> Path:
    from openpyxl import Workbook
    path = FIX / "listado_clientes_proveedores.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes y Proveedores"
    ws.append(["Tipo", "Nombre", "CIF / DNI", "Sector / Puesto"])
    rows = [
        ("Cliente", "Distribuciones García S.L.", "B11111111", "Distribución"),
        ("Cliente", "Mercadona S.A.", "A22222222", "Retail"),
        ("Cliente", "Supermercados Día", "A33333333", "Retail"),
        ("Proveedor", "Suministros Dentales Ibéricos S.L.", "B44444444", "Material dental"),
        ("Proveedor", "Material Médico Europa S.A.", "A55555555", "Material médico"),
        ("Banco", "Banco Santander", "A66666666", "Entidad financiera"),
        ("Banco", "CaixaBank", "A77777777", "Entidad financiera"),
        ("Empleado", "Laura Gómez Ruiz", "12345678Z", "Directora comercial"),
        ("Empleado", "Javier Torres Molina", "87654321X", "Responsable de compras"),
        ("Sociedad", "Inversiones Halcón S.L.", "B88888888", "Holding"),
        ("Sociedad", "Patrimonial Ramírez S.A.", "A99999999", "Patrimonial"),
    ]
    for r in rows:
        ws.append(r)
    wb.save(str(path))
    return path


def make_pptx() -> Path:
    from pptx import Presentation
    path = FIX / "presentacion_proyecto_halcon.pptx"
    prs = Presentation()
    slides_text = [
        ("Proyecto Halcón", "Resumen de la operación — Confidencial"),
        ("Partes de la operación",
         "Vendedor: Grupo Ramírez S.L.\nComprador: Salud Capital Partners\n"
         "Fondo coinversor: Mediterráneo Capital Partners\n"
         "Asesor financiero: Implica Corporate Finance\nDespacho legal: Garrigues"),
        ("Dirección", "CEO y fundador: Alberto Ramírez\nLa familia controla el 100% del capital."),
    ]
    blank = prs.slide_layouts[1]
    for title, body in slides_text:
        s = prs.slides.add_slide(blank)
        s.shapes.title.text = title
        s.placeholders[1].text = body
    prs.save(str(path))
    return path


def detect_entities(path: Path):
    from implica_anon import formats
    from implica_anon.detectors import detect_candidates, cluster_variants
    t0 = time.time()
    texts = formats.extract_text(path)
    candidates = detect_candidates(texts, skip_ner=False)
    clusters = cluster_variants(candidates)
    elapsed = time.time() - t0
    name_kinds = {"ORG", "PER", "CLIENTE", "PROVEEDOR", "DEUDOR", "GRUPO", "BANCO", "DOMINIO"}
    ents = {c.canonical for c in clusters if c.kind in name_kinds}
    return ents, elapsed


def _norm(s: str) -> str:
    return s.strip().lower()


def _hits(detected: set, truth: list) -> list:
    dl = [_norm(d) for d in detected]
    return [t for t in truth if any(_norm(t) in d or d in _norm(t) for d in dl)]


def _false_positives(detected: set, truth: list) -> list:
    tl = [_norm(t) for t in truth]
    return [d for d in detected if not any(_norm(d) in t or t in _norm(d) for t in tl)]


def _spacy_ok() -> bool:
    try:
        from implica_anon.detectors import _load_nlp
        return _load_nlp() is not None
    except Exception:
        return False


def collect():
    paths = {"teaser": make_teaser_pdf(), "excel": make_excel(), "pptx": make_pptx()}
    from implica_anon import gliner_detector
    gliner_active = gliner_detector.is_enabled()
    spacy_ok = _spacy_ok()

    docs = []
    for k, p in paths.items():
        truth = GROUND_TRUTH[k]
        os.environ.pop("IMPLICA_ENABLE_GLINER", None)
        base_ents, base_t = detect_entities(p)
        if gliner_active:
            os.environ["IMPLICA_ENABLE_GLINER"] = "true"
            gln_ents, gln_t = detect_entities(p)
        else:
            gln_ents, gln_t = base_ents, base_t

        base_hits = _hits(base_ents, truth)
        gln_hits = _hits(gln_ents, truth)
        new_useful = [e for e in gln_hits if e not in base_hits]
        rec = _doc_reco(gliner_active, len(base_hits), len(gln_hits),
                        len(_false_positives(base_ents, truth)),
                        len(_false_positives(gln_ents, truth)))
        docs.append({
            "doc": k, "label": DOC_LABEL[k], "file": p.name, "truth": len(truth),
            "base_hits": len(base_hits), "gln_hits": len(gln_hits),
            "new_useful": new_useful,
            "base_fp": len(_false_positives(base_ents, truth)),
            "gln_fp": len(_false_positives(gln_ents, truth)),
            "base_t": round(base_t, 2), "gln_t": round(gln_t, 2),
            "recommendation": rec,
        })
    os.environ.pop("IMPLICA_ENABLE_GLINER", None)

    totals = {kk: sum(d[kk] for d in docs) for kk in
              ("truth", "base_hits", "gln_hits", "base_fp", "gln_fp")}
    totals["new_useful"] = sum(len(d["new_useful"]) for d in docs)
    totals["base_t"] = round(sum(d["base_t"] for d in docs), 2)
    totals["gln_t"] = round(sum(d["gln_t"] for d in docs), 2)

    return {
        "meta": {
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "gliner_active": gliner_active,
            "spacy_ok": spacy_ok,
            "model": os.environ.get("IMPLICA_GLINER_MODEL", gliner_detector.DEFAULT_MODEL),
        },
        "docs": docs,
        "totals": totals,
    }


def _doc_reco(active, bh, gh, bfp, gfp):
    if not active:
        return "Pendiente (GLiNER no activo)"
    extra = gh - bh
    if extra >= 2 and (gfp - bfp) <= extra:
        return "Activar"
    if extra >= 1:
        return "Opcional"
    return "No aporta"


def _overall_reco(res):
    if not res["meta"]["gliner_active"]:
        return "Seguir probando"
    t = res["totals"]
    extra = t["gln_hits"] - t["base_hits"]
    extra_fp = t["gln_fp"] - t["base_fp"]
    # Si solo mejora texto libre, recomendamos activación selectiva
    narrative_gain = sum(d["gln_hits"] - d["base_hits"] for d in res["docs"] if d["doc"] in ("teaser", "pptx"))
    if extra >= 3 and extra_fp <= extra and narrative_gain >= extra * 0.6:
        return "Activar GLiNER solo para documentos narrativos"
    if extra >= 4 and extra_fp <= extra:
        return "Activar GLiNER para todo"
    if extra >= 1:
        return "Seguir probando"
    return "Mantener GLiNER apagado"


# --- Secciones estáticas (no dependen de los números) ---

SECCION_EQUIPO = """## 5. Pensado para equipo

Si esto lo usara **todo el equipo de Implica**, hay dos formas de desplegarlo:

| | A) App centralizada (servidor/Azure) | B) App local por analista |
|---|---|---|
| Instalación | 1 sola (la hace IT) | Cada persona instala Python/Torch/modelo |
| Acceso | Por navegador | App en cada equipo |
| Mantenimiento | IT controla deps, modelo y versiones | Cada uno mantiene lo suyo |
| Consistencia | Todos la misma versión/modelo | Riesgo de versiones distintas |
| Confidencialidad | En el tenant de Implica (ya privado) | En el equipo de cada analista |
| Coste de arranque | Imagen más pesada (Torch+modelo) una vez | Torch+modelo en N equipos |
| Facilidad para el equipo | Alta | Baja |

**Recomendación para Implica: opción A (centralizada).**
- El anonimizador YA está desplegado centralizado (Azure Container App en el tenant de Implica). Añadir GLiNER ahí (cuando se apruebe) es **una sola decisión**, no diez instalaciones.
- Nadie instala Torch (pesado) ni descarga modelos; el equipo solo usa el navegador.
- IT controla el modelo, su **licencia** y la versión → consistencia y trazabilidad.
- La ventaja de privacidad de la opción B es **redundante**: la app central ya vive dentro del perímetro de Implica (igual que OneDrive). No se gana privacidad instalando en cada equipo, solo complejidad.
"""

SECCION_REQUISITOS = """## 6. Requisitos antes de que lo use el equipo

- [ ] **spaCy funcionando** (es_core_news_md) en el servidor — desbloquear DLLs en Application Control si aplica.
- [ ] **GLiNER instalado solo si IT lo aprueba** (no por defecto).
- [ ] **Modelo descargado manualmente** una vez (`python -m implica_anon.gliner_detector`); después, offline.
- [ ] **Licencia documentada** del modelo (Apache-2.0 verificada para `urchade/gliner_multi-v2.1`).
- [ ] **Pruebas con documentos ficticios** (este benchmark).
- [ ] **Pruebas con documentos anonimizados** reales (sin PII) para validar en casos propios.
- [ ] **Decisión de IT antes de producción**.
- [ ] **No activar en Azure todavía** (la imagen no lleva GLiNER; sigue dormido).
"""


def write_reports(res):
    (OUT / "benchmark_resultado.md").write_text(_build_md(res), encoding="utf-8")
    (OUT / "benchmark_resultado.html").write_text(_build_html(res), encoding="utf-8")

    with (OUT / "benchmark_resultado.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Documento", "Entidades reales", "Aciertos sin GLiNER",
                    "Aciertos con GLiNER", "Nuevas utiles", "FP sin", "FP con",
                    "Tiempo sin (s)", "Tiempo con (s)", "Recomendacion"])
        for d in res["docs"]:
            w.writerow([d["label"], d["truth"], d["base_hits"], d["gln_hits"],
                        len(d["new_useful"]), d["base_fp"], d["gln_fp"],
                        d["base_t"], d["gln_t"], d["recommendation"]])
        t = res["totals"]
        w.writerow(["TOTAL", t["truth"], t["base_hits"], t["gln_hits"], t["new_useful"],
                    t["base_fp"], t["gln_fp"], t["base_t"], t["gln_t"], _overall_reco(res)])

    (OUT / "benchmark_resultado.json").write_text(
        json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")


def _build_html(res):
    import html as _h
    m, t = res["meta"], res["totals"]
    active = m["gliner_active"]
    overall = _overall_reco(res)

    def esc(x):
        return _h.escape(str(x))

    warns = ""
    if not m["spacy_ok"]:
        warns += ('<div class="warn">⚠️ spaCy no estaba operativo: el motor clásico cayó a '
                  'heurísticas. Los números infravaloran el clásico — repite en una máquina '
                  'con spaCy.</div>')
    if not active:
        warns += ('<div class="warn">⚠️ GLiNER no estaba activo: la columna «con GLiNER» = '
                  'clásico. Repite con <code>IMPLICA_ENABLE_GLINER=true</code> y GLiNER '
                  'instalado para la comparación real.</div>')

    if not active:
        resumen = (f"No se ejecutó la comparación real con GLiNER (estaba desactivado). El motor "
                   f"clásico detectó <b>{t['base_hits']}/{t['truth']}</b> entidades con "
                   f"<b>{t['base_fp']} falsos positivos</b> en {t['base_t']:.1f}s. Para decidir si "
                   f"GLiNER aporta hay que repetir con GLiNER instalado. Recomendación: "
                   f"<b>seguir probando</b>.")
    else:
        extra = t["gln_hits"] - t["base_hits"]
        extra_fp = t["gln_fp"] - t["base_fp"]
        resumen = (f"Sobre {t['truth']} entidades reales: clásico {t['base_hits']}, con GLiNER "
                   f"{t['gln_hits']} (<b>{extra:+d}</b> útiles). Falsos positivos "
                   f"{t['base_fp']}→{t['gln_fp']} (<b>{extra_fp:+d}</b>). Tiempo "
                   f"{t['base_t']:.1f}s→{t['gln_t']:.1f}s. Recomendación: <b>{esc(overall)}</b>.")

    rows = ""
    for d in res["docs"]:
        rows += (f"<tr><td>{esc(d['label'])}</td><td>{d['truth']}</td>"
                 f"<td>{d['base_hits']}</td><td class='hl'>{d['gln_hits']}</td>"
                 f"<td>{len(d['new_useful'])}</td><td>{d['base_fp']}</td><td>{d['gln_fp']}</td>"
                 f"<td>{d['base_t']:.2f}s</td><td>{d['gln_t']:.2f}s</td>"
                 f"<td>{esc(d['recommendation'])}</td></tr>")
    rows += (f"<tr class='tot'><td>TOTAL</td><td>{t['truth']}</td><td>{t['base_hits']}</td>"
             f"<td class='hl'>{t['gln_hits']}</td><td>{t['new_useful']}</td><td>{t['base_fp']}</td>"
             f"<td>{t['gln_fp']}</td><td>{t['base_t']:.2f}s</td><td>{t['gln_t']:.2f}s</td>"
             f"<td><b>{esc(overall)}</b></td></tr>")

    def fverdict(doc_key):
        if not active:
            return "Pendiente de medir"
        d = next((x for x in res["docs"] if x["doc"] == doc_key), None)
        if not d:
            return "Inferido"
        g = d["gln_hits"] - d["base_hits"]
        return f"Sí aporta (+{g})" if g > 0 else "No aporta"

    func_rows = (
        f"<tr><td>Teasers</td><td>{fverdict('teaser')}</td><td>Medido (texto libre, su punto fuerte)</td></tr>"
        f"<tr><td>PDFs con texto</td><td>{fverdict('teaser')}</td><td>Misma ruta NER que el teaser</td></tr>"
        f"<tr><td>PowerPoints</td><td>{fverdict('pptx')}</td><td>Medido</td></tr>"
        f"<tr><td>Word</td><td>Inferido = igual que teaser</td><td>Misma ruta NER (no medido aparte)</td></tr>"
        f"<tr><td>Excels CONTABLES (PGC)</td><td>No aplica</td><td>El motor PGC ya es 100%; GLiNER no interviene</td></tr>"
        f"<tr><td>Excels NO contables</td><td>{fverdict('excel')}</td><td>Medido (listado de nombres)</td></tr>"
    )

    opts = ["Activar GLiNER para todo", "Activar GLiNER solo para documentos narrativos",
            "Mantener GLiNER apagado", "Seguir probando"]
    opts_html = " · ".join(
        (f"<b class='pill'>{esc(o)}</b>" if o == overall else f"<span class='opt'>{esc(o)}</span>")
        for o in opts)

    reqs = "".join(f"<li>{esc(x)}</li>" for x in [
        "spaCy funcionando (es_core_news_md) en el servidor",
        "GLiNER instalado solo si IT lo aprueba (no por defecto)",
        "Modelo descargado manualmente una vez; después offline",
        "Licencia documentada (Apache-2.0 para urchade/gliner_multi-v2.1)",
        "Pruebas con documentos ficticios (este benchmark)",
        "Pruebas con documentos anonimizados reales (sin PII)",
        "Decisión de IT antes de producción",
        "No activar en Azure todavía (sigue dormido)",
    ])

    return f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Benchmark GLiNER — Implica</title>
<style>
:root{{--teal:#003E51;--teal2:#00BBB4;--lime:#C2D500;--orange:#FF6B00;--ink:#1b2733;--bg:#f4f6f7;}}
*{{box-sizing:border-box}}
body{{font-family:'Segoe UI',Helvetica,Arial,sans-serif;margin:0;background:var(--bg);color:var(--ink);line-height:1.5}}
.wrap{{max-width:980px;margin:0 auto;padding:0 24px 64px}}
header{{background:var(--teal);color:#fff;padding:32px 24px;border-bottom:5px solid var(--lime)}}
header .wrap{{padding-bottom:0}}
h1{{margin:0;font-size:26px}} header .meta{{opacity:.8;font-size:13px;margin-top:6px}}
h2{{color:var(--teal);border-left:5px solid var(--teal2);padding-left:12px;margin-top:38px}}
.warn{{background:#fff4e6;border:1px solid var(--orange);color:#8a4b00;padding:10px 14px;border-radius:8px;margin:14px 0;font-size:14px}}
.card{{background:#fff;border:1px solid #e2e8ea;border-radius:12px;padding:18px 22px;margin-top:14px;box-shadow:0 1px 3px rgba(0,0,0,.04)}}
table{{width:100%;border-collapse:collapse;margin-top:12px;font-size:14px;background:#fff;border-radius:10px;overflow:hidden}}
th{{background:var(--teal);color:#fff;text-align:left;padding:10px}}
td{{padding:9px 10px;border-top:1px solid #eef1f2}}
tr.tot td{{font-weight:700;background:#eef7f7;border-top:2px solid var(--teal2)}}
td.hl{{background:rgba(0,187,180,.10);font-weight:600}}
.reco{{background:var(--teal);color:#fff;padding:18px 22px;border-radius:12px;font-size:18px;margin-top:14px}}
.reco b.pill,b.pill{{background:var(--lime);color:var(--teal);padding:2px 10px;border-radius:20px}}
.opt{{opacity:.7;font-size:14px}}
ul{{line-height:1.9}} code{{background:#eef1f2;padding:1px 5px;border-radius:4px;font-size:13px}}
footer{{margin-top:40px;font-size:12px;opacity:.6;text-align:center}}
.rec-A{{border-left:5px solid var(--lime)}}
</style></head>
<body>
<header><div class="wrap"><h1>🔒 Benchmark GLiNER — Implica Anonimizador</h1>
<div class="meta">Generado: {esc(m['timestamp'])} · Modelo: {esc(m['model'])} · GLiNER activo: {("sí" if active else "no")}</div></div></header>
<div class="wrap">
{warns}

<h2>1. Resumen ejecutivo</h2>
<div class="card">{resumen}</div>

<h2>2. Tabla comparativa</h2>
<table>
<tr><th>Documento</th><th>Reales</th><th>Sin GLiNER</th><th>Con GLiNER</th><th>Nuevas útiles</th>
<th>FP sin</th><th>FP con</th><th>Tiempo sin</th><th>Tiempo con</th><th>Recomendación</th></tr>
{rows}
</table>

<h2>3. Conclusión funcional (por tipo de documento)</h2>
<table>
<tr><th>Tipo</th><th>¿Merece la pena GLiNER?</th><th>Nota</th></tr>
{func_rows}
</table>

<h2>4. Recomendación de activación</h2>
<div class="reco">➡️ <b class="pill">{esc(overall)}</b></div>
<p style="margin-top:10px">Opciones consideradas: {opts_html}</p>

<h2>5. Pensado para equipo</h2>
<div class="card rec-A">
<p>Si lo usara <b>todo el equipo</b>, hay dos formas de desplegarlo:</p>
<table>
<tr><th></th><th>A) Centralizada (servidor/Azure)</th><th>B) Local por analista</th></tr>
<tr><td>Instalación</td><td>1 sola (IT)</td><td>Cada persona instala todo</td></tr>
<tr><td>Acceso</td><td>Por navegador</td><td>App en cada equipo</td></tr>
<tr><td>Mantenimiento</td><td>IT controla deps/modelo/versión</td><td>Cada uno el suyo</td></tr>
<tr><td>Consistencia</td><td>Todos igual</td><td>Riesgo de versiones distintas</td></tr>
<tr><td>Facilidad equipo</td><td>Alta</td><td>Baja</td></tr>
</table>
<p style="margin-top:12px"><b style="color:var(--teal)">Recomendación para Implica: opción A (centralizada).</b>
El anonimizador ya está centralizado (Azure en el tenant de Implica). Añadir GLiNER ahí es
<b>una decisión, no diez instalaciones</b>; nadie instala Torch ni descarga modelos; IT controla
modelo, licencia y versión. La privacidad de B es redundante: la app central ya vive dentro del
perímetro de Implica.</p>
</div>

<h2>6. Requisitos antes de que lo use el equipo</h2>
<div class="card"><ul>{reqs}</ul></div>

<h2>7. Archivos de esta carpeta</h2>
<div class="card"><ul>
<li><code>benchmark_resultado.html</code> — este informe</li>
<li><code>benchmark_resultado.md</code> · <code>.csv</code> · <code>.json</code></li>
<li>Documentos ficticios usados: teaser (PDF), Excel y PPT</li>
</ul></div>

<footer>Implica Corporate Finance · Benchmark local · GLiNER desactivado por defecto · datos ficticios</footer>
</div></body></html>"""


def _build_md(res):
    m, t = res["meta"], res["totals"]
    active, spok = m["gliner_active"], m["spacy_ok"]
    overall = _overall_reco(res)

    L = []
    L.append("# Benchmark GLiNER — Implica Anonimizador\n")
    L.append(f"_Generado: {m['timestamp']} · Modelo: `{m['model']}`_\n")
    if not spok:
        L.append("> ⚠️ **spaCy no estaba operativo**: el motor clásico cayó a heurísticas. "
                 "Los números infravaloran el clásico. Repite en una máquina con spaCy.\n")
    if not active:
        L.append("> ⚠️ **GLiNER no estaba activo**: la columna 'con GLiNER' = clásico. "
                 "Repite con `IMPLICA_ENABLE_GLINER=true` y GLiNER instalado.\n")

    # 1. Resumen ejecutivo
    L.append("## 1. Resumen ejecutivo\n")
    if not active:
        L.append("No se ha ejecutado la comparación real con GLiNER (estaba desactivado). "
                 "El motor clásico detectó **{}/{}** entidades con **{} falsos positivos** en "
                 "{:.1f}s. Para decidir si GLiNER aporta, hay que repetir el benchmark con "
                 "GLiNER instalado y activo en una máquina con spaCy operativo. Hasta entonces, "
                 "la recomendación es **seguir probando**.\n".format(
                     t["base_hits"], t["truth"], t["base_fp"], t["base_t"]))
    else:
        extra = t["gln_hits"] - t["base_hits"]
        extra_fp = t["gln_fp"] - t["base_fp"]
        L.append("Sobre {} entidades reales, el motor clásico detectó {} y con GLiNER {} "
                 "(**{:+d}** entidades útiles). Falsos positivos: {} → {} (**{:+d}**). "
                 "Tiempo total: {:.1f}s → {:.1f}s. GLiNER {} en documentos narrativos "
                 "(teaser/PPT) y no interviene en Excel contable PGC. "
                 "Recomendación global: **{}**.\n".format(
                     t["truth"], t["base_hits"], t["gln_hits"], extra,
                     t["base_fp"], t["gln_fp"], extra_fp, t["base_t"], t["gln_t"],
                     "aporta" if extra > 0 else "no aporta valor neto", overall))

    # 2. Tabla comparativa
    L.append("## 2. Tabla comparativa\n")
    L.append("| Documento | Reales | Sin GLiNER | Con GLiNER | Nuevas útiles | FP sin | FP con | Tiempo sin | Tiempo con | Recomendación |")
    L.append("|---|---|---|---|---|---|---|---|---|---|")
    for d in res["docs"]:
        L.append("| {} | {} | {} | {} | {} | {} | {} | {:.2f}s | {:.2f}s | {} |".format(
            d["label"], d["truth"], d["base_hits"], d["gln_hits"], len(d["new_useful"]),
            d["base_fp"], d["gln_fp"], d["base_t"], d["gln_t"], d["recommendation"]))
    L.append("| **TOTAL** | {} | {} | {} | {} | {} | {} | {:.2f}s | {:.2f}s | **{}** |".format(
        t["truth"], t["base_hits"], t["gln_hits"], t["new_useful"], t["base_fp"],
        t["gln_fp"], t["base_t"], t["gln_t"], overall))
    L.append("")
    new_all = [e for d in res["docs"] for e in d["new_useful"]]
    if new_all:
        L.append("**Entidades nuevas útiles que añadió GLiNER:** " + ", ".join(new_all) + "\n")

    # 3. Conclusión funcional
    L.append("## 3. Conclusión funcional (por tipo de documento)\n")
    def verdict(doc_key, measured=True):
        if not active:
            return "Pendiente de medir"
        d = next((x for x in res["docs"] if x["doc"] == doc_key), None)
        if d is None:
            return "Inferido"
        g = d["gln_hits"] - d["base_hits"]
        return "Sí aporta (+{})".format(g) if g > 0 else "No aporta"
    L.append("| Tipo | ¿Merece la pena GLiNER? | Nota |")
    L.append("|---|---|---|")
    L.append(f"| Teasers | {verdict('teaser')} | Medido (texto libre, es su punto fuerte) |")
    L.append(f"| PDFs con texto | {verdict('teaser')} | Misma ruta NER que el teaser |")
    L.append(f"| PowerPoints | {verdict('pptx')} | Medido |")
    L.append("| Word | Inferido = igual que teaser | Misma ruta NER (no medido aparte) |")
    L.append("| Excels CONTABLES (PGC) | No aplica | El motor PGC ya es 100%; GLiNER no interviene por diseño |")
    L.append(f"| Excels NO contables | {verdict('excel')} | Medido (listado de nombres) |")
    L.append("")

    # 4. Recomendación de activación
    L.append("## 4. Recomendación de activación\n")
    L.append(f"**➡️ {overall}**\n")
    opts = ["Activar GLiNER para todo", "Activar GLiNER solo para documentos narrativos",
            "Mantener GLiNER apagado", "Seguir probando"]
    L.append("Opciones consideradas: " + " · ".join(
        ("**" + o + "**" if o == overall else o) for o in opts) + "\n")

    # 5 y 6 (estáticas)
    L.append(SECCION_EQUIPO)
    L.append(SECCION_REQUISITOS)

    # 7. Output adicional
    L.append("## 7. Output adicional\n")
    L.append("- `benchmark_resultado.csv` — la tabla para abrir en Excel.")
    L.append("- `benchmark_resultado.json` — datos completos para trazabilidad.\n")
    return "\n".join(L)


def run():
    print("Generando fixtures y ejecutando benchmark...")
    res = collect()
    write_reports(res)
    t = res["totals"]
    print(f"\nTotal: clásico {t['base_hits']}/{t['truth']} (FP {t['base_fp']}) · "
          f"GLiNER {t['gln_hits']}/{t['truth']} (FP {t['gln_fp']}) · recomendación: {_overall_reco(res)}")
    print(f"\nInformes generados en: {OUT}")
    print("  - benchmark_resultado.html  (ábrelo con doble clic / abrir_benchmark.bat)")
    print("  - benchmark_resultado.md / .csv / .json")
    print("  - documentos ficticios: teaser (PDF), Excel, PPT")
    if not res["meta"]["gliner_active"]:
        print("\n⚠️ GLiNER no estaba activo → repite con IMPLICA_ENABLE_GLINER=true "
              "e IMPLICA_GLINER_ALLOW_DOWNLOAD=true para la comparación real.")


if __name__ == "__main__":
    run()
