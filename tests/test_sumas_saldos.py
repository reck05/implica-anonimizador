"""Diagnóstico de la herramienta actual contra un sumas y saldos grande.

Mide:
- Tiempo de extracción + detección
- Cuántos candidatos detecta vs. los reales en el fixture
- Calidad: falsos positivos (cuentas genéricas detectadas como empresa) y
  falsos negativos (clientes que no detecta)
"""
import io
import re
import sys
import time
from collections import Counter
from pathlib import Path

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", line_buffering=True)
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", line_buffering=True)

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from implica_anon import formats
from implica_anon.detectors import cluster_variants, detect_candidates

FIXTURE = Path(__file__).parent / "fixtures" / "sumas_saldos_sintetico.xlsx"


def get_ground_truth() -> dict[str, set[str]]:
    """Lee el fixture y extrae los nombres que SABEMOS son entidades a anonimizar."""
    wb = load_workbook(str(FIXTURE), data_only=True, read_only=True)
    ws = wb["Sumas y Saldos 2024"]
    truth: dict[str, set[str]] = {
        "clientes": set(),
        "proveedores": set(),
        "deudores": set(),
        "personal": set(),
        "bancos": set(),
        "generic_accounts": set(),
    }
    for row in ws.iter_rows(min_row=6, values_only=True):
        if not row[0]:
            continue
        code = str(row[0])
        desc = str(row[1] or "")
        if code.startswith("43"):
            # Limpiar paréntesis con CIF
            nombre = re.sub(r"\s*\([^)]*\)", "", desc).strip()
            truth["clientes"].add(nombre)
        elif code.startswith(("410", "411")):
            nombre = re.sub(r",\s*CIF.*$", "", desc).strip()
            truth["proveedores"].add(nombre)
        elif code.startswith("440"):
            truth["deudores"].add(desc.strip())
        elif code.startswith("460"):
            # "Anticipo Juan Pérez García"
            nombre = desc.replace("Anticipo", "").strip()
            truth["personal"].add(nombre)
        elif code.startswith("572"):
            # "Banco Santander - ES..."
            nombre = desc.split(" - ")[0].strip()
            truth["bancos"].add(nombre)
        elif len(code) <= 4:
            # Cuentas PGC genéricas — NO deberían detectarse como empresa
            truth["generic_accounts"].add(desc.strip())
    wb.close()
    return truth


def main():
    print(f"=== Diagnóstico contra {FIXTURE.name} ===\n")

    truth = get_ground_truth()
    print("Ground truth del fixture:")
    print(f"  Clientes esperados:        {len(truth['clientes'])}")
    print(f"  Proveedores esperados:     {len(truth['proveedores'])}")
    print(f"  Deudores esperados:        {len(truth['deudores'])}")
    print(f"  Personal esperado:         {len(truth['personal'])}")
    print(f"  Bancos esperados:          {len(truth['bancos'])}")
    print(f"  Cuentas PGC genéricas:     {len(truth['generic_accounts'])} (NO deberían anonimizarse)")
    print()

    # --- Extracción ---
    t0 = time.time()
    texts = formats.extract_text(FIXTURE)
    t1 = time.time()
    print(f"Extracción: {len(texts)} fragmentos en {t1-t0:.1f}s")

    # --- Detección ---
    print("Ejecutando detección (spaCy + regex)...")
    t2 = time.time()
    candidates = detect_candidates(texts)
    t3 = time.time()
    print(f"Detección: {len(candidates)} candidatos en {t3-t2:.1f}s")

    # --- Clustering ---
    t4 = time.time()
    clusters = cluster_variants(candidates)
    t5 = time.time()
    print(f"Clustering: {len(clusters)} clusters en {t5-t4:.1f}s")

    print(f"\nTIEMPO TOTAL: {t5-t0:.1f}s")
    print()

    # Stats por tipo
    by_kind: dict[str, list] = {}
    for c in clusters:
        by_kind.setdefault(c.kind, []).append(c)
    for kind, items in by_kind.items():
        print(f"  [{kind}] {len(items)} clusters")

    # --- Calidad: cobertura de clientes ---
    detected_orgs = set()
    for c in clusters:
        if c.kind == "ORG":
            for v in c.variants:
                detected_orgs.add(v)

    expected_clients = truth["clientes"] | truth["proveedores"]
    matched = 0
    for client in expected_clients:
        # Match si cualquier variante del cliente está en lo detectado
        for v in detected_orgs:
            if client.lower() in v.lower() or v.lower() in client.lower():
                matched += 1
                break
    coverage = matched / len(expected_clients) * 100 if expected_clients else 0
    print(f"\nCobertura clientes+proveedores: {matched}/{len(expected_clients)} ({coverage:.1f}%)")

    # --- Falsos positivos: cuentas PGC genéricas detectadas como ORG/PER ---
    fp_generic = []
    for c in clusters:
        if c.kind in ("ORG", "PER"):
            for v in c.variants:
                if v.strip() in truth["generic_accounts"]:
                    fp_generic.append(v)
    if fp_generic:
        print(f"\n⚠ Falsos positivos (cuentas PGC genéricas detectadas como entidad):")
        for v in fp_generic[:10]:
            print(f"    - {v}")
        if len(fp_generic) > 10:
            print(f"    ... y {len(fp_generic)-10} más")
    else:
        print("\n✓ Sin falsos positivos en cuentas genéricas")

    # --- Bancos: ¿detectó Santander, BBVA, etc.? ---
    detected_banks = []
    for c in clusters:
        if c.kind == "ORG":
            for v in c.variants:
                if any(b in v for b in ["Santander", "BBVA", "CaixaBank", "Sabadell"]):
                    detected_banks.append(v)
                    break
    print(f"\nBancos detectados como ORG: {len(detected_banks)}")
    for b in detected_banks[:5]:
        print(f"    - {b}")

    # --- Muestra de clusters ORG con variantes ---
    multi_variant_orgs = [c for c in clusters if c.kind == "ORG" and len(c.variants) > 1]
    print(f"\nClusters ORG con múltiples variantes: {len(multi_variant_orgs)}")
    for c in multi_variant_orgs[:5]:
        print(f"    - canónico: '{c.canonical}'")
        print(f"      variantes: {c.variants[:5]}{'...' if len(c.variants) > 5 else ''}")

    # --- Sample de clientes NO detectados ---
    not_detected = []
    for client in list(expected_clients)[:50]:
        found = False
        for v in detected_orgs:
            if client.lower() in v.lower() or v.lower() in client.lower():
                found = True
                break
        if not found:
            not_detected.append(client)
    if not_detected:
        print(f"\nClientes/proveedores no detectados (muestra):")
        for n in not_detected[:10]:
            print(f"    - {n}")


if __name__ == "__main__":
    main()
